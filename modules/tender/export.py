from __future__ import annotations

from datetime import datetime, timezone
from typing import List, Optional

from docx import Document
from docx.shared import Pt

from modules.tender.schemas import (
    CoverageSummary,
    EvidenceGap,
    TenderExtraction,
    TenderMatrix,
)

_TENDER_TEXT_MAX_CHARS = 5000


def export_win_pack_docx(
    out_path: str,
    tender_text: str,
    extraction: TenderExtraction,
    matrix: TenderMatrix,
    summary: Optional[CoverageSummary],
    gaps: Optional[List[EvidenceGap]],
) -> str:
    doc = Document()

    # A) Title + timestamp
    doc.add_heading("Tender Win Pack", level=0)
    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    doc.add_paragraph(f"Generated: {ts}")

    # B) Coverage summary
    if summary is not None:
        doc.add_heading("Coverage Summary", level=1)
        doc.add_paragraph(
            f"Total: {summary.total}  |  Met: {summary.met}  |  "
            f"Partial: {summary.partial}  |  Missing: {summary.missing}  |  "
            f"Score: {summary.score:.2%}"
        )
        if summary.missing_ids:
            doc.add_paragraph("Missing requirement IDs: " + ", ".join(summary.missing_ids))

    # C) Compliance matrix table
    doc.add_heading("Compliance Matrix", level=1)
    req_map = {r.id: r for r in extraction.requirements}
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Requirement ID"
    hdr[1].text = "Category"
    hdr[2].text = "Status"
    hdr[3].text = "Notes"
    for cell in hdr:
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True

    for row in matrix.rows:
        req = req_map.get(row.requirement_id)
        category = req.category if req else ""
        cells = table.add_row().cells
        cells[0].text = row.requirement_id
        cells[1].text = category
        cells[2].text = row.status
        cells[3].text = row.notes

    # D) Missing evidence
    if gaps:
        doc.add_heading("Missing Evidence", level=1)
        for gap in gaps:
            doc.add_paragraph(
                f"{gap.requirement_id}: " + ", ".join(gap.missing_evidence),
                style="List Bullet",
            )

    # E) Original tender text (truncated)
    doc.add_heading("Original Tender Text", level=1)
    truncated = tender_text[:_TENDER_TEXT_MAX_CHARS]
    if len(tender_text) > _TENDER_TEXT_MAX_CHARS:
        truncated += "\n[truncated]"
    doc.add_paragraph(truncated)

    doc.save(out_path)
    return out_path
